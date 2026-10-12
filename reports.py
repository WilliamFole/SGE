# -*- coding: utf-8 -*-
import io
import os
import re
import base64
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

from services import arredondar_moeda

BRAND_COLOR = colors.HexColor("#2F5233")
BRAND_COLOR_HEX = "2F5233"


def _caminho_logo_empresa():
    from database import get_config, get_base_dir
    logo_cfg = get_config("logo_empresa", "")
    if not logo_cfg:
        return None
    if logo_cfg.startswith("data:"):
        return None
    caminho = os.path.join(get_base_dir(), "static", "img", logo_cfg)
    if os.path.exists(caminho):
        return caminho
    caminho2 = os.path.join(get_base_dir(), logo_cfg)
    if os.path.exists(caminho2):
        return caminho2
    return None


def _logo_data_url_bytes():
    from database import get_config, get_base_dir
    logo_cfg = get_config("logo_empresa", "")
    if not logo_cfg:
        return None
    if logo_cfg.startswith("data:image") and ";base64," in logo_cfg:
        try:
            header, data = logo_cfg.split(",", 1)
            ext = "png" if "png" in header else "jpg"
            return base64.b64decode(data), ext
        except Exception:
            return None
    caminho = os.path.join(get_base_dir(), "static", "img", logo_cfg)
    if not os.path.exists(caminho):
        caminho = os.path.join(get_base_dir(), logo_cfg)
    if os.path.exists(caminho):
        ext = caminho.rsplit(".", 1)[-1].lower()
        with open(caminho, "rb") as f:
            return f.read(), ext
    return None


def _dados_empresa():
    from database import get_company_info
    return get_company_info()


def _cabecalho_empresa_texto():
    emp = _dados_empresa()
    linhas = []
    if emp.get("nome"):
        linhas.append(emp["nome"])
    morada_nuit_contacto = []
    if emp.get("morada"):
        morada_nuit_contacto.append(emp["morada"])
    if emp.get("nuit"):
        morada_nuit_contacto.append(f"NUIT: {emp['nuit']}")
    if emp.get("contacto"):
        morada_nuit_contacto.append(f"Tel.: {emp['contacto']}")
    if emp.get("email"):
        morada_nuit_contacto.append(emp["email"])
    if morada_nuit_contacto:
        linhas.append("  |  ".join(morada_nuit_contacto))
    return linhas


def formatar_moeda_br(valor):
    """Formata valor com duas casas decimais e separador de milhar (ex: 1.000,00)."""
    if valor is None:
        return "0,00"
    try:
        return f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "0,00"


def formatar_inteiro(valor):
    """Formata inteiro com separador de milhar (ex: 1.000)."""
    if valor is None:
        return "0"
    try:
        val = int(float(valor))
        return f"{val:,}".replace(",", ".")
    except (TypeError, ValueError):
        return "0"


def gerar_pdf(titulo, colunas, linhas, subtitulo=None, empresa_titulo=None, is_cotacao=False, landscape_mode=False):
    buf = io.BytesIO()
    if not landscape_mode and not is_cotacao:
        landscape_mode = len(colunas) > 6 or any(len(str(c)) > 30 for c in colunas)
    pagesize = landscape(A4) if landscape_mode or is_cotacao else A4
    doc = SimpleDocTemplate(buf, pagesize=pagesize,
                             topMargin=1.2*cm, bottomMargin=1.2*cm,
                             leftMargin=1.2*cm, rightMargin=1.2*cm)

    styles = getSampleStyleSheet()
    # ---- Estilos ALINHADOS À ESQUERDA (canto superior esquerdo) ----
    style_titulo_empresa_esq = ParagraphStyle(
        "EmpresaTituloEsq", parent=styles["Title"],
        fontSize=15, textColor=BRAND_COLOR,
        alignment=TA_LEFT, spaceAfter=2, spaceBefore=0,
    )
    style_sub_empresa_esq = ParagraphStyle(
        "EmpresaSubEsq", parent=styles["Normal"],
        fontSize=9, textColor=colors.HexColor("#4B4B4B"),
        alignment=TA_LEFT, spaceAfter=1, spaceBefore=0,
    )
    elementos = []

    emp = _dados_empresa()
    cabecalho_linhas = _cabecalho_empresa_texto()

    # ---- Logo (à direita, opcional) ----
    logo_caminho = _caminho_logo_empresa()
    img = None
    if logo_caminho:
        try:
            img = RLImage(logo_caminho, width=2.8*cm, height=2.8*cm)
            img.hAlign = "RIGHT"
        except Exception:
            img = None
    else:
        logo_bytes = _logo_data_url_bytes()
        if logo_bytes:
            try:
                from io import BytesIO
                b, ext = logo_bytes
                bio = BytesIO(b)
                img = RLImage(bio, width=2.8*cm, height=2.8*cm)
                img.hAlign = "RIGHT"
            except Exception:
                img = None

    titulo_emp = empresa_titulo or emp.get("nome") or "SGE — Sistema de Gestão de Estoque"
    sub = cabecalho_linhas[1:] if len(cabecalho_linhas) > 1 else []

    # ---- Coluna esquerda: INFORMAÇÃO DA EMPRESA (top-left) ----
    texto_empresa = [Paragraph(titulo_emp, style_titulo_empresa_esq)]
    for s in sub:
        texto_empresa.append(Paragraph(s, style_sub_empresa_esq))

    # ---- Coluna direita: LOGO ----
    col_dir = [img] if img else [Paragraph("", styles["Normal"])]

    # Larguras: info à esquerda ocupa a maior parte, logo à direita
    largura_logo = 3.2 * cm
    largura_info = doc.width - largura_logo

    cabecalho_tabela = Table([[texto_empresa, col_dir]],
                             colWidths=[largura_info, largura_logo])
    cabecalho_tabela.setStyle(TableStyle([
        ("VALIGN", (0,0), (0,0), "TOP"),
        ("VALIGN", (1,0), (1,0), "TOP"),
        ("ALIGN", (0,0), (0,0), "LEFT"),
        ("ALIGN", (1,0), (1,0), "RIGHT"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
        ("TOPPADDING", (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2),
    ]))
    elementos.append(cabecalho_tabela)

    # Linha separadora fina
    from reportlab.platypus import HRFlowable
    elementos.append(Spacer(1, 0.15*cm))
    elementos.append(HRFlowable(width="100%", thickness=0.6,
                                 color=BRAND_COLOR, spaceAfter=6))
    elementos.append(Spacer(1, 0.2*cm))

    # ---- Título do relatório (centrado) ----
    if emp.get("nome"):
        elementos.append(Paragraph(titulo, ParagraphStyle(
            "TituloRel", parent=styles["Heading2"], alignment=TA_CENTER,
            textColor=BRAND_COLOR, spaceAfter=2,
        )))
    else:
        elementos.append(Paragraph(titulo, styles["Title"]))

    if subtitulo:
        elementos.append(Paragraph(subtitulo, ParagraphStyle(
            "SubRel", parent=styles["Normal"], alignment=TA_CENTER, spaceAfter=2,
        )))
    elementos.append(Paragraph(
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        ParagraphStyle("DataGer", parent=styles["Normal"],
                       alignment=TA_CENTER, fontSize=8, textColor=colors.grey),
    ))
    elementos.append(Spacer(1, 0.5*cm))

    # ---- Formatação dos dados (mesma lógica de antes) ----
    COLUNAS_PRE_FORMATADAS = {
        "qtd. orig.", "qtd. base", "saldo acum. (un.)",
        "saldo inicial (qtd)", "entradas (qtd)", "saídas (qtd)",
        "saldo final (qtd)",
    }
    dados_formatados = [colunas]
    for linha in linhas:
        linha_formatada = []
        for idx, val in enumerate(linha):
            if colunas[idx].strip().lower() in COLUNAS_PRE_FORMATADAS:
                linha_formatada.append(str(val))
            col_nome = colunas[idx].lower()
            num_val = None
            if isinstance(val, (int, float)):
                num_val = float(val)
            elif isinstance(val, str):
                texto = val.strip().replace(' ', '')
                if ',' in texto and '.' in texto:
                    texto_norm = texto.replace('.', '').replace(',', '.')
                elif ',' in texto and '.' not in texto:
                    texto_norm = texto.replace(',', '.')
                else:
                    texto_norm = texto
                if re.fullmatch(r"-?\d+(?:\.\d+)?", texto_norm):
                    try:
                        num_val = float(texto_norm)
                    except Exception:
                        num_val = None
            if num_val is not None:
                if any(p in col_nome for p in ("preço","preco","valor","total","lucro","custo","venda","receita")):
                    linha_formatada.append(formatar_moeda_br(num_val))
                elif any(p in col_nome for p in ("qtd","quantidade","saldo","estoque","unid","numero","número","unidade","minimo","maximo")):
                    linha_formatada.append(formatar_inteiro(num_val))
                else:
                    if abs(num_val - round(num_val)) < 1e-9:
                        linha_formatada.append(formatar_inteiro(num_val))
                    else:
                        linha_formatada.append(formatar_moeda_br(num_val))
            else:
                linha_formatada.append(str(val))
        dados_formatados.append(linha_formatada)

    tabela = Table(dados_formatados, repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), BRAND_COLOR),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("FONTNAME", (0,0), (-1,-1), "Times-Roman"),
        ("FONTNAME", (0,0), (-1,0), "Times-Bold"),
        ("GRID", (0,0), (-1,-1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ALIGN", (0,0), (-1,-1), "LEFT"),
        ("ALIGN", (1,1), (-1,-1), "RIGHT"),
    ])) 
    elementos.append(tabela)
    doc.build(elementos)
    buf.seek(0)
    return buf

def gerar_excel(titulo, colunas, linhas, empresa_titulo=None):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31] if titulo else "Relatorio"

    emp = _dados_empresa()
    nome_emp = empresa_titulo or emp.get("nome") or "SGE — Sistema de Gestão de Estoque"
    linha_atual = 1

    logo_info = _logo_data_url_bytes()
    if logo_info:
        try:
            import tempfile
            dados_bytes, ext = logo_info
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
            tmp.write(dados_bytes)
            tmp.close()
            img = XLImage(tmp.name)
            img.width = 80
            img.height = 80
            ws.add_image(img, "A1")
            linha_atual = 6
        except Exception:
            pass

    header_fill = PatternFill(start_color=BRAND_COLOR_HEX, end_color=BRAND_COLOR_HEX,
                              fill_type="solid")
    header_font = Font(name='Times New Roman', bold=True, color="FFFFFF", size=12)
    default_font = Font(name='Times New Roman', size=11)
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    ws.cell(row=linha_atual, column=1, value=nome_emp).font = Font(
        name='Times New Roman', bold=True, color=BRAND_COLOR_HEX, size=14)
    ws.merge_cells(start_row=linha_atual, start_column=1,
                   end_row=linha_atual, end_column=max(3, len(colunas)))
    linha_atual += 1

    sub_linhas = []
    if emp.get("morada"):
        sub_linhas.append(emp["morada"])
    if emp.get("nuit"):
        sub_linhas.append(f"NUIT: {emp['nuit']}")
    if emp.get("contacto"):
        sub_linhas.append(f"Tel.: {emp['contacto']}")
    if emp.get("email"):
        sub_linhas.append(emp["email"])
    if sub_linhas:
        ws.cell(row=linha_atual, column=1, value="  |  ".join(sub_linhas)).font = Font(
            size=10, color="555555")
        ws.merge_cells(start_row=linha_atual, start_column=1,
                       end_row=linha_atual, end_column=max(3, len(colunas)))
        linha_atual += 1

    ws.cell(row=linha_atual, column=1, value=titulo).font = Font(
        bold=True, color=BRAND_COLOR_HEX, size=12)
    ws.merge_cells(start_row=linha_atual, start_column=1,
                   end_row=linha_atual, end_column=max(3, len(colunas)))
    linha_atual += 1
    ws.cell(row=linha_atual, column=1,
            value=f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = Font(
        size=9, color="666666")
    linha_atual += 2

    header_row = linha_atual
    for i, col in enumerate(colunas, 1):
        c = ws.cell(row=header_row, column=i, value=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    for idx, linha in enumerate(linhas, 1):
        for j, val in enumerate(linha, 1):
            c = ws.cell(row=header_row + idx, column=j, value=val)
            c.border = thin_border
            c.font = default_font
            if idx % 2 == 0:
                c.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            col_nome = colunas[j-1].lower()
            if any(p in col_nome for p in ("preço", "preco", "valor", "total", "lucro", "custo", "venda", "receita")):
                if isinstance(val, (int, float)):
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal="right")
            elif isinstance(val, (int, float)):
                c.number_format = '#,##0'
                c.alignment = Alignment(horizontal="right")

    for col_cells in ws.columns:
        largura = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells) + 2
        first_cell = next((c for c in col_cells if hasattr(c, 'column_letter')), None)
        if first_cell is not None:
            ws.column_dimensions[first_cell.column_letter].width = min(largura, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def gerar_relatorio_precos(conn):
    produtos = conn.execute(
        """SELECT p.codigo, p.nome, c.nome AS categoria, p.preco_venda
           FROM produtos p LEFT JOIN categorias c ON c.id_categoria=p.id_categoria
           WHERE p.ativo=1 ORDER BY p.nome"""
    ).fetchall()
    colunas = ["Código", "Produto", "Categoria", "Preço de Venda (MZN)"]
    linhas = []
    for p in produtos:
        linhas.append([p["codigo"], p["nome"], p["categoria"] or "—", arredondar_moeda(p["preco_venda"])])
    return colunas, linhas


def gerar_pdf_cotacao(cotacao, itens, empresa_titulo=None):
    """Cotação em A4 retrato, layout tipo fatura."""
    from reportlab.platypus import HRFlowable
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=1.0*cm, bottomMargin=1.0*cm,
        leftMargin=1.2*cm, rightMargin=1.2*cm,
    )

    styles = getSampleStyleSheet()
    emp = _dados_empresa()

    s_emp = ParagraphStyle("emp", parent=styles["Normal"], fontSize=11,
                           textColor=BRAND_COLOR, alignment=TA_LEFT, leading=13)
    s_emp_sub = ParagraphStyle("empsub", parent=styles["Normal"], fontSize=8,
                                textColor=colors.HexColor("#555555"),
                                alignment=TA_LEFT, leading=10)
    s_doc_titulo = ParagraphStyle("dt", parent=styles["Heading1"], fontSize=18,
                                   textColor=BRAND_COLOR, alignment=TA_RIGHT,
                                   spaceAfter=2)
    s_doc_sub = ParagraphStyle("ds", parent=styles["Normal"], fontSize=9,
                                alignment=TA_RIGHT, textColor=colors.HexColor("#444444"))
    s_secao = ParagraphStyle("sec", parent=styles["Normal"], fontSize=9,
                              textColor=colors.white, fontName="Times-Bold")
    s_cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8, leading=10)

    elementos = []

    # ---------- Cabeçalho ----------
    logo_cell = ""
    p = _caminho_logo_empresa()
    if p:
        try:
            logo_cell = RLImage(p, width=2.4*cm, height=2.4*cm)
        except Exception:
            logo_cell = ""

    emp_txt = [Paragraph(f"<b>{emp.get('nome') or 'SGE'}</b>", s_emp)]
    linhas_emp = []
    if emp.get("morada"):   linhas_emp.append(emp["morada"])
    if emp.get("nuit"):     linhas_emp.append(f"NUIT: {emp['nuit']}")
    if emp.get("contacto"): linhas_emp.append(f"Tel.: {emp['contacto']}")
    if emp.get("email"):    linhas_emp.append(emp["email"])
    for l in linhas_emp:
        emp_txt.append(Paragraph(l, s_emp_sub))
    if not linhas_emp:
        emp_txt.append(Paragraph("—", s_emp_sub))

    valor_total_txt = formatar_moeda_br(cotacao["valor_total"])
    doc_txt = [
        Paragraph("COTAÇÃO", s_doc_titulo),
        Paragraph(f"<b>Nº:</b> {cotacao['numero']}", s_doc_sub),
        Paragraph(f"<b>Emissão:</b> {cotacao['data_emissao'][:10]}", s_doc_sub),
        Paragraph(f"<b>Validade:</b> {cotacao['data_validade'][:10]}", s_doc_sub),
        Paragraph(f"<b>Status:</b> {cotacao['status']}", s_doc_sub),
        Paragraph(f"<b>Total:</b> {valor_total_txt}", s_doc_sub),
    ]

    header_tbl = Table(
        [[logo_cell, emp_txt, doc_txt]],
        colWidths=[2.6*cm, doc.width - 2.6*cm - 5.5*cm, 5.5*cm],
    )
    header_tbl.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (2,0), (2,0), "RIGHT"),
        ("LEFTPADDING", (0,0), (-1,-1), 2),
        ("RIGHTPADDING", (0,0), (-1,-1), 2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    elementos.append(header_tbl)
    elementos.append(HRFlowable(width="100%", thickness=1.0,
                                 color=BRAND_COLOR, spaceAfter=6))

    # ---------- Bloco cliente ----------
    cliente_cell = [
        Paragraph("<b>CLIENTE</b>", s_secao),
        Paragraph(f"<b>{cotacao['nome_cliente'] or '—'}</b>", s_cell),
    ]
    if cotacao.get("email_cliente"):
        cliente_cell.append(Paragraph(f"Email: {cotacao['email_cliente']}", s_cell))
    if cotacao.get("telefone_cliente"):
        cliente_cell.append(Paragraph(f"Tel.: {cotacao['telefone_cliente']}", s_cell))

    obs_cell = [
        Paragraph("<b>OBSERVAÇÕES</b>", s_secao),
        Paragraph(cotacao.get("observacoes") or "—", s_cell),
    ]

    info_tbl = Table([[cliente_cell, obs_cell]], colWidths=[doc.width*0.5, doc.width*0.5])
    info_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F2F2F2")),
        ("BOX", (0,0), (-1,-1), 0.4, colors.grey),
        ("INNERGRID", (0,0), (-1,-1), 0.4, colors.grey),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 6),
        ("RIGHTPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    elementos.append(info_tbl)
    elementos.append(Spacer(1, 0.4*cm))

    # ---------- Tabela de itens ----------
    cab = ["#", "Código", "Descrição", "Qtd.", "Un.", "Preço Unit.", "Desc.%", "Subtotal"]
    dados = [cab]
    for i, it in enumerate(itens, 1):
        dados.append([
            str(i),
            it["codigo"],
            it["produto_nome"],
            str(int(it["quantidade"])),
            "Un",
            formatar_moeda_br(it["preco_unitario"]),
            f"{it['desconto_percentual']:.0f}%",
            formatar_moeda_br(it["valor_total"]),
        ])
    dados.append(["", "", "", "", "", "", "TOTAL",
                  formatar_moeda_br(cotacao["valor_total"])])

    larguras = [
        doc.width*0.04, doc.width*0.12, doc.width*0.34,
        doc.width*0.08, doc.width*0.06, doc.width*0.12,
        doc.width*0.08, doc.width*0.16,
    ]
    tbl = Table(dados, colWidths=larguras, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), BRAND_COLOR),
        ("TEXTCOLOR",  (0,0), (-1,0), colors.white),
        ("FONTNAME",   (0,0), (-1,0), "Times-Bold"),
        ("FONTNAME",   (0,1), (-1,-2), "Times-Roman"),
        ("FONTNAME",   (0,-1), (-1,-1), "Times-Bold"),
        ("FONTSIZE",   (0,0), (-1,-1), 8),
        ("GRID",       (0,0), (-1,-1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0,1), (-1,-2), [colors.white, colors.HexColor("#F7F7F7")]),
        ("ALIGN", (3,1), (-1,-1), "RIGHT"),
        ("ALIGN", (0,0), (0,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("SPAN", (0,-1), (6,-1)),
        ("BACKGROUND", (0,-1), (-1,-1), colors.HexColor("#EAF1E8")),
        ("TEXTCOLOR", (0,-1), (-1,-1), BRAND_COLOR),
    ]))
    elementos.append(tbl)

    elementos.append(Spacer(1, 0.5*cm))
    elementos.append(HRFlowable(width="100%", thickness=0.4,
                                 color=colors.grey, spaceAfter=4))
    elementos.append(Paragraph(
        f"Cotação gerada em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
        f"Validade: {cotacao['data_validade'][:10]} · "
        f"Este documento não constitui fatura.",
        ParagraphStyle("rod", parent=styles["Normal"], fontSize=7,
                       textColor=colors.grey, alignment=TA_CENTER),
    ))

    doc.build(elementos)
    buf.seek(0)
    return buf



# ---------------------------------------------------------------------------
# Fatura e Recibo (vendas a crédito)
# ---------------------------------------------------------------------------
def _cabecalho_fatura(doc_obj, titulo, subtitulo_linhas):
    styles = getSampleStyleSheet()
    emp = _dados_empresa()
    s_emp = ParagraphStyle("emp2", parent=styles["Normal"], fontSize=10,
                           textColor=BRAND_COLOR, leading=12)
    s_emp_sub = ParagraphStyle("empsub2", parent=styles["Normal"], fontSize=8,
                               textColor=colors.HexColor("#555"), leading=10)
    s_tit = ParagraphStyle("tit2", parent=styles["Heading1"], fontSize=16,
                           alignment=TA_RIGHT, textColor=BRAND_COLOR)
    s_sub = ParagraphStyle("sub2", parent=styles["Normal"], fontSize=9,
                           alignment=TA_RIGHT, textColor=colors.HexColor("#333"))

    logo_cell = ""
    p = _caminho_logo_empresa()
    if p:
        try: logo_cell = RLImage(p, width=2.2*cm, height=2.2*cm)
        except Exception: pass

    esq = [Paragraph(f"<b>{emp.get('nome') or 'SGE'}</b>", s_emp)]
    for l in filter(None, [
        emp.get("morada"),
        f"NUIT: {emp['nuit']}" if emp.get("nuit") else None,
        f"Tel.: {emp['contacto']}" if emp.get("contacto") else None,
        emp.get("email"),
    ]):
        esq.append(Paragraph(l, s_emp_sub))

    dir_ = [Paragraph(titulo, s_tit)]
    for l in subtitulo_linhas:
        dir_.append(Paragraph(l, s_sub))

    tbl = Table([[logo_cell, esq, dir_]],
                colWidths=[2.4*cm, 9.2*cm, doc_obj.width - 2.4*cm - 9.2*cm])
    tbl.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("ALIGN",(2,0),(2,0),"RIGHT"),
        ("LEFTPADDING",(0,0),(-1,-1),2),
        ("RIGHTPADDING",(0,0),(-1,-1),2),
    ]))
    return tbl


def _logo_banco_path():
    from database import get_config, get_base_dir
    nome = get_config("banco_logo", "")
    if not nome:
        return None
    p = os.path.join(get_base_dir(), "static", "img", nome)
    return p if os.path.exists(p) else None


def gerar_pdf_fatura(conta, itens=None):
    """Fatura A4 retrato — layout ERP (mockup aprovado)."""
    from reportlab.platypus import HRFlowable, PageBreak
    from database import get_conn as _gc

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=0.9*cm, bottomMargin=0.9*cm,
        leftMargin=1.0*cm, rightMargin=1.0*cm,
    )
    styles = getSampleStyleSheet()
    emp = _dados_empresa()

    # ---- Estilos ----
    s_lbl = ParagraphStyle("lbl", parent=styles["Normal"], fontSize=6.5,
                           textColor=colors.HexColor("#555"), leading=8)
    s_val = ParagraphStyle("val", parent=styles["Normal"], fontSize=8, leading=10)
    s_val_b = ParagraphStyle("valb", parent=styles["Normal"], fontSize=8,
                             leading=10, fontName="Times-Bold")
    s_tit = ParagraphStyle("tit", parent=styles["Heading2"], fontSize=13,
                           textColor=colors.black, leading=15,
                           fontName="Times-Bold")
    s_orig = ParagraphStyle("orig", parent=styles["Normal"], fontSize=8,
                            alignment=TA_RIGHT, textColor=colors.HexColor("#333"))
    s_sec = ParagraphStyle("sec", parent=styles["Normal"], fontSize=8,
                           textColor=colors.white, fontName="Times-Bold",
                           leading=10)
    s_rod = ParagraphStyle("rd", parent=styles["Normal"], fontSize=6.5,
                           alignment=TA_CENTER, textColor=colors.grey, leading=8)

    elementos = []

    # ============================================================
    # BLOCO 1 — Cabeçalho
    # ============================================================
    # Coluna ESQUERDA: logo + empresa
    empresa_col = []
    p_logo = _caminho_logo_empresa()
    if p_logo:
        try:
            img = RLImage(p_logo, width=3.2*cm, height=3.2*cm)
            img.hAlign = "LEFT"
            empresa_col.append(img)
        except Exception:
            pass
    empresa_col.append(Paragraph(f"<b>{emp.get('nome') or 'SGE'}</b>",
                                  ParagraphStyle("en", parent=styles["Normal"],
                                                 fontSize=9, leading=11,
                                                 textColor=colors.HexColor("#333"))))
    def _kv(lbl, val):
        if not val:
            return None
        return Paragraph(
            f"<font size=6.5 color='#555'>{lbl}:</font> "
            f"<font size=8>{val}</font>", s_val)
    for lbl, val in [
        ("NUIT N.º", emp.get("nuit")),
        ("", emp.get("morada")),
        ("", emp.get("email")),
        ("Tel.", emp.get("contacto")),
    ]:
        par = _kv(lbl, val)
        if par: empresa_col.append(par)

    # Coluna DIREITA: dados do cliente
    cliente_col = [Paragraph("Exmo.(s) Sr.(s)", s_val)]
    cliente_col.append(Paragraph(" ",
                                  ParagraphStyle("sep", parent=styles["Normal"],
                                                 fontSize=8, textColor=colors.HexColor("#999"))))
    cliente_col.append(Paragraph(f"<b>{(conta.get('nome_cliente') or '—').upper()}</b>", s_val_b))
    if conta.get("nif_cliente"):
        cliente_col.append(Paragraph(f"NUIT: {conta['nif_cliente']}", s_val))
    if conta.get("endereco_cliente"):
        cliente_col.append(Paragraph(conta["endereco_cliente"], s_val))
    if conta.get("localidade_cliente"):
        cliente_col.append(Paragraph(conta["localidade_cliente"], s_val))
    if conta.get("pais_cliente"):
        cliente_col.append(Paragraph(conta["pais_cliente"], s_val))

    hdr = Table(
        [[empresa_col, cliente_col]],
        colWidths=[doc.width*0.55, doc.width*0.45],
    )
    hdr.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (1,0), (1,0), "LEFT"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
    ]))
    elementos.append(hdr)

    # Pág. X/Y (topo direito) — usa canvas para desenhar
    elementos.append(Spacer(1, 0.15*cm))
    elementos.append(HRFlowable(width="100%", thickness=0.8,
                                 color=colors.grey, spaceBefore=2, spaceAfter=6))

    # ============================================================
    # BLOCO 2 — Título do documento
    # ============================================================
    numero = conta.get("numero_fatura") or f"FT 2026/{conta['id_conta']:05d}"
    tipo = emp.get("tipo_documento") or "Factura"
    titulo = Paragraph(f"{tipo} &nbsp; {numero}", s_tit)
    origem = Paragraph("Original", s_orig)
    tbl_tit = Table([[titulo, origem]], colWidths=[doc.width*0.75, doc.width*0.25])
    tbl_tit.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "BOTTOM"),
        ("ALIGN", (1,0), (1,0), "RIGHT"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
        ("LINEBELOW", (0,0), (-1,-1), 0.4, colors.grey),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    elementos.append(tbl_tit)
    elementos.append(Spacer(1, 0.15*cm))

    # Linha "Condições:"
    cond_label = "crédito" if (conta.get("condicao_pagamento") or "").lower() != "pronto pagamento" else "pronto pagamento"
    elementos.append(Paragraph(f"Condições: <b>{cond_label}</b>", s_val))
    elementos.append(Spacer(1, 0.25*cm))

    # ============================================================
    # BLOCO 3 — Linha de metadados (data / venc / cond / câmbio / moeda)
    # ============================================================
    taxa_iva = float(conta.get("taxa_iva") or emp.get("taxa_iva_padrao") or 16)
    isento = bool(conta.get("isento_iva"))

    meta_cab = ["Data", "Vencim.", "Cond. Pag.", "Desc.C.", "Desc.F.",
                "Câmbio", "Moeda", "V/Lei Contra"]
    meta_val = [
        conta["data_emissao"][:10],
        (conta.get("data_vencimento") or "—")[:10] if conta.get("data_vencimento") else "—",
        conta.get("condicao_pagamento") or emp.get("condicoes_pagamento_padrao") or "—",
        "0,00", "0,00",
        emp.get("cambio_padrao") or "0,00",
        emp.get("moeda") or "MZN",
        "",
    ]
    meta_tbl = Table([meta_cab, meta_val],
                     colWidths=[doc.width*0.12, doc.width*0.12, doc.width*0.15,
                                doc.width*0.10, doc.width*0.10, doc.width*0.12,
                                doc.width*0.10, doc.width*0.19])
    meta_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F2F2F2")),
        ("FONTSIZE", (0,0), (-1,-1), 7),
        ("FONTNAME", (0,0), (-1,0), "Times-Bold"),
        ("TEXTCOLOR", (0,0), (-1,0), colors.HexColor("#333")),
        ("GRID", (0,0), (-1,-1), 0.3, colors.grey),
        ("ALIGN", (3,1), (-1,-1), "RIGHT"),
        ("TOPPADDING", (0,0), (-1,-1), 2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2),
    ]))
    elementos.append(meta_tbl)
    elementos.append(Spacer(1, 0.3*cm))

    # ============================================================
    # BLOCO 4 — Itens
    # ============================================================
    if not itens:
        conn = _gc()
        if conta.get("id_movimento"):
            mov = conn.execute(
                """SELECT m.*, p.nome AS produto_nome, p.codigo AS produto_codigo
                   FROM movimentos m JOIN produtos p ON p.id_produto=m.id_produto
                   WHERE m.id_movimento=?""",
                (conta["id_movimento"],)
            ).fetchone()
            if mov:
                itens = [{
                    "produto_codigo": mov["produto_codigo"],
                    "produto_nome": mov["produto_nome"],
                    "unidade_original": mov["unidade_original"],
                    "quantidade_original": mov["quantidade_original"],
                    "preco_unitario": mov["preco_unitario"],
                    "valor_total": mov["valor_total"],
                }]
        conn.close()
    itens = itens or []

    cab_itens = ["Artigo", "Descrição", "Qtd.", "Un.",
                 "P. Unit.", "Desc.", "IVA", "Valor"]
    dados = [cab_itens]
    subtotal = 0.0
    iva_por_linha_total = 0.0
    for it in itens:
        u = (it.get("unidade_original") or "Unidade").strip().lower()
        q = float(it.get("quantidade_original") or 0)
        qtd_txt = f"{q:.3f}".replace(".", ",") if u == "kg" else f"{int(round(q))}"
        total = float(it.get("valor_total") or 0)
        subtotal += total
        iva_txt = "Isento" if isento else f"{taxa_iva:.0f}%"
        dados.append([
            it.get("produto_codigo", "—"),
            it.get("produto_nome", "—"),
            qtd_txt,
            it.get("unidade_original", "UN"),
            formatar_moeda_br(it.get("preco_unitario")),
            "0,00",
            iva_txt,
            formatar_moeda_br(total),
        ])

    larguras = [
        doc.width*0.11, doc.width*0.35, doc.width*0.075, doc.width*0.055,
        doc.width*0.11, doc.width*0.08, doc.width*0.06, doc.width*0.16,
    ]
    tbl_itens = Table(dados, colWidths=larguras, repeatRows=1)
    tbl_itens.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2F5233")),
        ("TEXTCOLOR",  (0,0), (-1,0), colors.white),
        ("FONTNAME",   (0,0), (-1,0), "Times-Bold"),
        ("FONTSIZE",   (0,0), (-1,-1), 7.5),
        ("GRID",       (0,0), (-1,-1), 0.3, colors.grey),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F7F7F7")]),
        ("ALIGN", (2,1), (2,-1), "RIGHT"),
        ("ALIGN", (4,1), (-1,-1), "RIGHT"),
        ("ALIGN", (6,1), (6,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2),
    ]))
    elementos.append(tbl_itens)
    elementos.append(Spacer(1, 0.4*cm))

    # ============================================================
    # BLOCO 5 — Total + Resumo de Impostos
    # ============================================================
    total_com_iva = float(conta["valor_total"] or 0)
    if isento or taxa_iva <= 0:
        iliquido = total_com_iva
        iva_val = 0.0
    else:
        iliquido = total_com_iva / (1 + taxa_iva/100)
        iva_val = total_com_iva - iliquido

    # --- Coluna esquerda: quadro resumo de impostos ---
    q_cab = ["Taxa", "Incidência", "Total", "Motivo Isenção"]
    if isento:
        q_lin1 = ["Isento", formatar_moeda_br(iliquido), "0,00",
                  emp.get("motivo_isencao_iva") or "Isento por lei"]
        q_lin2 = ["—", "0,00", "0,00", ""]
    else:
        q_lin1 = [f"IVA {taxa_iva:.0f}%", formatar_moeda_br(iliquido),
                  formatar_moeda_br(iva_val), emp.get("motivo_isencao_iva") or ""]
        q_lin2 = ["Isento", "0,00", "0,00", ""]

    q_tbl = Table([q_cab, q_lin1, q_lin2],
                  colWidths=[1.6*cm, 2.2*cm, 2.2*cm, doc.width*0.20])
    q_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F2F2F2")),
        ("FONTSIZE", (0,0), (-1,-1), 7),
        ("FONTNAME", (0,0), (-1,0), "Times-Bold"),
        ("GRID", (0,0), (-1,-1), 0.3, colors.grey),
        ("ALIGN", (1,1), (2,-1), "RIGHT"),
        ("TOPPADDING", (0,0), (-1,-1), 2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2),
    ]))

    # --- Coluna direita: totais ---
    tot_linhas = [
        ["Mercadorias/Serviços",  formatar_moeda_br(iliquido)],
        ["Descontos Comerciais",  "-0,00"],
        ["Desconto Financeiro",   "0,00"],
        ["Portes",                "0,00"],
        ["Outros Serviços",       "0,00"],
        ["Adiantamentos",         "0,00"],
        ["IVA",                   formatar_moeda_br(iva_val)],
    ]
    tot_tbl = Table(tot_linhas, colWidths=[doc.width*0.22, doc.width*0.18],
                    hAlign="RIGHT")
    tot_tbl.setStyle(TableStyle([
        ("FONTSIZE", (0,0), (-1,-1), 7.5),
        ("ALIGN", (1,0), (1,-1), "RIGHT"),
        ("LINEABOVE", (0,-1), (-1,-1), 0.3, colors.grey),
        ("TOPPADDING", (0,0), (-1,-1), 2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2),
    ]))

    # --- Bloco "Total [MZN]" (destaque) ---
    total_lbl = Paragraph(
        f"<font size=10><b>Total [{emp.get('moeda') or 'MZN'}]</b></font>",
        ParagraphStyle("tl", parent=styles["Normal"], alignment=TA_RIGHT))
    total_num = Paragraph(
        f"<font size=12 color='#2F5233'><b>{formatar_moeda_br(total_com_iva)}</b></font>",
        ParagraphStyle("tn", parent=styles["Normal"], alignment=TA_RIGHT))
    total_tbl = Table([[total_lbl, total_num]],
                      colWidths=[doc.width*0.22, doc.width*0.18], hAlign="RIGHT")
    total_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#EAF1E8")),
        ("BOX", (0,0), (-1,-1), 0.4, colors.HexColor("#2F5233")),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
        ("LEFTPADDING", (0,0), (-1,-1), 4),
        ("RIGHTPADDING", (0,0), (-1,-1), 4),
    ]))

    # Combinar: tabela 2 colunas
    dir_col = [tot_tbl, total_tbl]
    linha_resumo = Table([[q_tbl, dir_col]],
                         colWidths=[doc.width*0.50, doc.width*0.50])
    linha_resumo.setStyle(TableStyle([
        ("VALIGN", (0,0), (0,0), "TOP"),
        ("VALIGN", (1,0), (1,0), "TOP"),
        ("ALIGN", (1,0), (1,0), "RIGHT"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
        ("TOPPADDING", (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
    ]))
    elementos.append(linha_resumo)
    elementos.append(Spacer(1, 0.6*cm))

    # ============================================================
    # BLOCO 6 — Coordenadas bancárias + Texto legal
    # ============================================================
    elementos.append(HRFlowable(width="100%", thickness=0.4, color=colors.grey,
                                 spaceBefore=4, spaceAfter=4))

    banco_col = [Paragraph("<b>Coordenadas Bancárias</b>",
                            ParagraphStyle("cb", parent=styles["Normal"],
                                           fontSize=9, leading=11))]
    banco_col.append(Paragraph(
        f"Beneficiário: <b>{emp.get('nome') or '—'}</b>", s_val))
    p_banco = _logo_banco_path()
    if p_banco:
        try:
            banco_col.append(Spacer(1, 0.15*cm))
            img_b = RLImage(p_banco, width=2.2*cm, height=1.2*cm)
            img_b.hAlign = "LEFT"
            banco_col.append(img_b)
        except Exception:
            pass
    elif emp.get("banco_nome"):
        banco_col.append(Paragraph(f"Banco: <b>{emp['banco_nome']}</b>", s_val))
    if emp.get("numero_conta"):
        banco_col.append(Paragraph(f"Nº Conta: {emp['numero_conta']}", s_val))
    if emp.get("iban"):
        banco_col.append(Paragraph(f"IBAN: {emp['iban']}", s_val))
    if emp.get("nib"):
        banco_col.append(Paragraph(f"NIB: {emp['nib']}", s_val))

    texto_legal = emp.get("texto_rodape_fatura") or (
        "Os bens e/ou serviços foram colocados à disposição do adquirente "
        "na data de emissão deste documento, de acordo com o regime legal aplicável."
    )
    legal_col = [Paragraph(texto_legal,
                            ParagraphStyle("lg", parent=styles["Normal"],
                                           fontSize=7, leading=9,
                                           textColor=colors.HexColor("#555")))]

    rod_tbl = Table([[banco_col, legal_col]],
                    colWidths=[doc.width*0.5, doc.width*0.5])
    rod_tbl.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
    ]))
    elementos.append(rod_tbl)
    elementos.append(Spacer(1, 0.5*cm))

    # Rodapé final
    elementos.append(HRFlowable(width="100%", thickness=0.3, color=colors.grey,
                                 spaceBefore=4, spaceAfter=4))
    elementos.append(Paragraph(
        "documento processado por computador — SGE-series",
        s_rod))

    # ---- Callback para "Pág. X/Y" no canto superior direito ----
    from reportlab.pdfgen import canvas as _canvas

    def _on_page(canv, doc_obj):
        canv.saveState()
        canv.setFont("Times-Roman", 7)
        canv.setFillColor(colors.HexColor("#333"))
        canv.drawRightString(doc_obj.pagesize[0] - 1.0*cm,
                              doc_obj.pagesize[1] - 0.6*cm,
                              f"Pág. {canv.getPageNumber()}/{doc_obj.page}")
        canv.restoreState()

    doc.build(elementos, onFirstPage=_on_page, onLaterPages=_on_page)
    buf.seek(0)
    return buf
